import hashlib
import importlib.util
import io
import os
from pathlib import Path
import stat
import subprocess
import tempfile
import unittest
from unittest import mock
import zipfile


KIT = Path(__file__).resolve().parents[3]
SPEC = importlib.util.spec_from_file_location("metadata_cleaner", KIT / "server/guest-clean-files.py")
cleaner = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(cleaner)

try:
    from PIL import Image, PngImagePlugin
    from docx import Document
    from reportlab.pdfgen import canvas
    from libmat2 import parser_factory
    REAL_DEPENDENCIES = True
except ImportError:
    REAL_DEPENDENCIES = False


class MetadataSafetyTests(unittest.TestCase):
    def setUp(self):
        # main() deliberately sets the standalone cleaner's private umask.
        # Restore process state so unrelated permission tests keep their own setup.
        original_umask = os.umask(0o077)
        os.umask(original_umask)
        self.addCleanup(os.umask, original_umask)
        self.temporary = tempfile.TemporaryDirectory()
        self.home = Path(self.temporary.name)
        self.source = self.home / "private.png"
        self.source.write_bytes(b"not a real image")

    def tearDown(self):
        self.temporary.cleanup()

    def test_unsupported_symlink_special_and_oversized_are_rejected(self):
        unsupported = self.home / "secret.txt"
        unsupported.write_text("private")
        link = self.home / "link.png"
        link.symlink_to(self.source)
        fifo = self.home / "pipe.png"
        os.mkfifo(fifo)
        oversized = self.home / "large.png"
        with oversized.open("wb") as stream:
            stream.truncate(cleaner.MAX_INPUT + 1)
        for path in (unsupported, link, fifo, oversized):
            with self.subTest(path=path.name), self.assertRaises(cleaner.CleaningError):
                cleaner.clean_file(path, self.home)
        self.assertFalse((self.home / "OnionDesk-Clean").exists())

    def test_private_output_required_and_directory_symlink_rejected(self):
        output = self.home / "OnionDesk-Clean"
        output.mkdir(mode=0o755)
        with self.assertRaises(cleaner.CleaningError):
            cleaner.clean_file(self.source, self.home)
        output.rmdir()
        output.symlink_to(self.home, target_is_directory=True)
        with self.assertRaises(cleaner.CleaningError):
            cleaner.clean_file(self.source, self.home)

    def test_cleaning_failure_never_publishes_a_copy(self):
        with mock.patch.object(cleaner, "_run_worker", side_effect=cleaner.CleaningError("failed")):
            with self.assertRaises(cleaner.CleaningError):
                cleaner.clean_file(self.source, self.home)
        self.assertEqual(list((self.home / "OnionDesk-Clean").iterdir()), [])
        self.assertEqual(self.source.read_bytes(), b"not a real image")

    def test_raw_staging_is_never_in_the_sendable_output_folder(self):
        def worker(staged):
            self.assertEqual(list((self.home / 'OnionDesk-Clean').iterdir()), [])
            self.assertEqual(staged.read_bytes(), self.source.read_bytes())
            self.assertEqual(staged.parent.parent, (self.home / '.oniondesk-processing').resolve())
            result = staged.with_name('input.cleaned.png')
            result.write_bytes(b'fixture cleaned copy')
            return result
        with mock.patch.object(cleaner, '_run_worker', side_effect=worker):
            output = cleaner.clean_file(self.source, self.home)
        self.assertEqual(list(output.parent.iterdir()), [output])
        self.assertEqual(list((self.home / '.oniondesk-processing').iterdir()), [])
        self.assertEqual(stat.S_IMODE((self.home / '.oniondesk-processing').stat().st_mode), 0o700)

    def test_private_processing_directory_cannot_be_a_link(self):
        (self.home / '.oniondesk-processing').symlink_to(self.home, target_is_directory=True)
        with self.assertRaises(cleaner.CleaningError):
            cleaner.clean_file(self.source, self.home)
        self.assertEqual(list((self.home / 'OnionDesk-Clean').iterdir()), [])

    def test_remaining_metadata_is_failure(self):
        for mime, fields in (("image/jpeg", {"GPSLatitude": "private"}),
                             ("application/pdf", {"format": "PDF-1.5", "author": "private"}),
                             ("application/pdf", {"creation-date": 123}),
                             ("application/pdf", {"format": "PDF-1.6"})):
            with self.subTest(mime=mime, fields=fields), self.assertRaises(cleaner.CleaningError):
                cleaner._verify_meta(fields, mime)
        cleaner._verify_meta({"format": "PDF-1.5", "creation-date": -1, "mod-date": -1}, "application/pdf")

    def test_docx_rejects_zip_traversal_and_expansion(self):
        path = self.home / "unsafe.docx"
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("[Content_Types].xml", "")
            archive.writestr("word/document.xml", "")
            archive.writestr("../outside.txt", "private")
        with self.assertRaises(cleaner.CleaningError):
            cleaner._check_docx(path)
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", "")
            archive.writestr("word/document.xml", "0" * 1024)
        with mock.patch.object(cleaner, "MAX_OUTPUT", 100), self.assertRaises(cleaner.CleaningError):
            cleaner._check_docx(path)

    def test_worker_timeout_kills_process_group(self):
        staged = self.home / "input.png"
        staged.write_bytes(b"input")
        process = mock.Mock(pid=12345)
        process.wait.side_effect = [subprocess.TimeoutExpired("parser", 1), 0]
        with mock.patch.object(cleaner.subprocess, "Popen", return_value=process), \
                mock.patch.object(cleaner.os, "killpg") as kill:
            with self.assertRaisesRegex(cleaner.CleaningError, "лимит времени"):
                cleaner._run_worker(staged, timeout=1)
        kill.assert_called_once_with(12345, cleaner.signal.SIGKILL)

    def test_docx_external_links_are_rejected(self):
        path = self.home / "linked.docx"
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("[Content_Types].xml", "")
            archive.writestr("word/document.xml", "")
            archive.writestr("_rels/.rels", '<Relationships><Relationship TargetMode="External" Target="https://example.invalid/private"/></Relationships>')
        with self.assertRaisesRegex(cleaner.CleaningError, "внешними ссылками"):
            cleaner._check_docx(path)

    def _nested_zip(self):
        nested = io.BytesIO()
        with zipfile.ZipFile(nested, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("expanded.txt", b"0" * (2 * 1024 * 1024))
        return nested.getvalue()

    def _docx(self, path, member_name=None, content=b"", relationships=None):
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
            archive.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body/></w:document>')
            if member_name:
                archive.writestr(member_name, content)
            if relationships:
                archive.writestr("word/_rels/document.xml.rels", relationships)
        return path

    def _assert_rejected_before_mat2(self, path):
        factory = mock.Mock()
        module = mock.Mock(parser_factory=factory)
        with mock.patch.dict(cleaner.sys.modules, {"libmat2": module}), \
                mock.patch.object(cleaner.resource, "setrlimit"), \
                mock.patch.object(cleaner, "MAX_OUTPUT", 1024 * 1024):
            with self.assertRaisesRegex(cleaner.CleaningError, "вложенными"):
                cleaner._worker(path)
        factory.get_parser.assert_not_called()

    def test_docx_nested_office_zip_bomb_rejected_before_mat2(self):
        path = self._docx(self.home / "bomb.docx", "word/embeddings/Microsoft_Excel_Worksheet1.xlsx", self._nested_zip())
        # Outer contents are tiny; the nested file would expand beyond the test
        # budget. Neither the inner ZIP nor libmat2 is ever opened.
        with zipfile.ZipFile(path) as outer:
            self.assertLess(sum(member.file_size for member in outer.infolist()), 16 * 1024)
        self._assert_rejected_before_mat2(path)

    def test_docx_zip_disguised_as_image_or_prefixed_image_rejected_before_mat2(self):
        nested = self._nested_zip()
        for prefix in (b"", b"\x89PNG\r\n\x1a\nimage-prefix"):
            with self.subTest(prefix=prefix):
                path = self._docx(self.home / "disguised.docx", "word/media/picture.png", prefix + nested)
                self._assert_rejected_before_mat2(path)

    def test_docx_ole_and_package_relationships_rejected_before_mat2(self):
        for kind in ("oleObject", "package"):
            with self.subTest(kind=kind):
                relations = ('<Relationships><Relationship Target="media/picture.png" '
                             'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
                             + kind + '"/></Relationships>')
                path = self._docx(self.home / "object.docx", "word/media/picture.png", b"ordinary-content", relations)
                self._assert_rejected_before_mat2(path)

    def test_cleaned_docx_nested_zip_rejected_before_verifier(self):
        source = self._docx(self.home / "input.docx")
        output = self._docx(self.home / "input.cleaned.docx", "word/media/picture.png", self._nested_zip())
        parser = mock.Mock(output_filename=str(output))
        parser.remove_all.return_value = True
        factory = mock.Mock()
        factory.get_parser.return_value = (parser, cleaner.TYPES[".docx"][0])
        with mock.patch.dict(cleaner.sys.modules, {"libmat2": mock.Mock(parser_factory=factory)}), \
                mock.patch.object(cleaner.resource, "setrlimit"):
            with self.assertRaisesRegex(cleaner.CleaningError, "вложенными"):
                cleaner._worker(source)
        factory.get_parser.assert_called_once_with(str(source))
        parser.remove_all.assert_called_once()

    def test_gui_cancel_and_success_show_output_and_limitations(self):
        output = self.home / "OnionDesk-Clean/file-123.png"
        cancelled = subprocess.CompletedProcess([], 1, b"", b"")
        success = subprocess.CompletedProcess([], 0, os.fsencode(str(self.source)) + b"\n", b"")
        with mock.patch.object(cleaner.os, "geteuid", return_value=1000), \
                mock.patch.object(cleaner.sys, "argv", ["clean-files.py"]), \
                mock.patch.object(cleaner.subprocess, "run", side_effect=[cancelled, success]), \
                mock.patch.object(cleaner, "clean_file", return_value=output) as clean, \
                mock.patch.object(cleaner, "_dialog") as dialog:
            self.assertEqual(cleaner.main(), 0)
            clean.assert_not_called()
            self.assertEqual(cleaner.main(), 0)
            clean.assert_called_once_with(self.source)
            self.assertIn(str(output), dialog.call_args.args[1])
            self.assertIn(cleaner.LIMITATION, dialog.call_args.args[1])

    def test_runtime_refuses_root(self):
        with mock.patch.object(cleaner.os, "geteuid", return_value=0):
            self.assertEqual(cleaner.main(), 1)


@unittest.skipUnless(REAL_DEPENDENCIES and os.geteuid() != 0,
                     "Actual Debian mat2, Pillow, python-docx, reportlab and an unprivileged user required")
class ActualMat2Tests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.home = Path(self.temporary.name)

    def tearDown(self):
        self.temporary.cleanup()

    def _clean_and_check(self, source):
        before_bytes = source.read_bytes()
        before = source.stat()
        result = cleaner.clean_file(source, self.home)
        after = source.stat()
        self.assertEqual((before.st_mtime_ns, before.st_ctime_ns, before.st_atime_ns),
                         (after.st_mtime_ns, after.st_ctime_ns, after.st_atime_ns))
        self.assertEqual(hashlib.sha256(source.read_bytes()).digest(), hashlib.sha256(before_bytes).digest())
        self.assertNotIn("private", result.name)
        self.assertRegex(result.name, r"^file-[0-9a-f]{24}\.(jpg|png|pdf|docx)$")
        self.assertEqual(stat.S_IMODE(result.stat().st_mode), 0o600)
        self.assertEqual(stat.S_IMODE(result.parent.stat().st_mode), 0o700)
        self.assertEqual(result.stat().st_mtime, cleaner.NEUTRAL_TIME)
        parser, mime = parser_factory.get_parser(str(result))
        cleaner._verify_meta(parser.get_meta(), mime)
        self.assertEqual([p for p in result.parent.iterdir() if p.name.startswith(".processing-")], [])
        return result

    def test_jpeg_gps_camera_author_are_removed_and_image_opens(self):
        source = self.home / "private-photo.jpg"
        Image.new("RGB", (48, 32), "red").save(source)
        subprocess.run(["/usr/bin/exiftool", "-overwrite_original", "-GPSLatitude=55.75",
                        "-GPSLatitudeRef=N", "-GPSLongitude=37.61", "-GPSLongitudeRef=E",
                        "-Artist=Private Author", "-Make=Private Camera", str(source)],
                       check=True, stdout=subprocess.DEVNULL)
        parser, _ = parser_factory.get_parser(str(source))
        self.assertIn("GPSLatitude", parser.get_meta())
        result = self._clean_and_check(source)
        with Image.open(result) as image:
            self.assertEqual(image.size, (48, 32))
            self.assertFalse(image.getexif())

    def test_png_author_text_are_removed_and_pixels_preserved(self):
        source = self.home / "private-image.png"
        text = PngImagePlugin.PngInfo()
        text.add_text("Author", "Private Author")
        text.add_text("Description", "Private Location")
        Image.new("RGB", (48, 32), "blue").save(source, pnginfo=text)
        result = self._clean_and_check(source)
        with Image.open(source) as original, Image.open(result) as image:
            self.assertEqual(image.convert("RGB").tobytes(), original.convert("RGB").tobytes())
            self.assertNotIn("Author", image.info)
            self.assertNotIn("Description", image.info)

    def test_pdf_author_title_are_removed_and_page_opens(self):
        source = self.home / "private-document.pdf"
        document = canvas.Canvas(str(source))
        document.setAuthor("Private Author")
        document.setTitle("Private Title")
        document.drawString(40, 40, "Document content stays visible")
        document.save()
        result = self._clean_and_check(source)
        fields = subprocess.run(["/usr/bin/exiftool", "-j", str(result)], check=True,
                                capture_output=True, text=True).stdout
        self.assertNotIn("Private Author", fields)
        self.assertNotIn("Private Title", fields)
        self.assertNotIn('"Author"', fields)
        self.assertIn('"PageCount": 1', fields)

    def test_docx_creator_last_editor_and_embedded_image_are_cleaned(self):
        source = self.home / "private-letter.docx"
        image = self.home / "embedded.png"
        text = PngImagePlugin.PngInfo()
        text.add_text("Author", "Private Embedded Author")
        Image.new("RGB", (8, 8), "green").save(image, pnginfo=text)
        document = Document()
        document.core_properties.author = "Private Author"
        document.core_properties.last_modified_by = "Private Editor"
        document.core_properties.title = "Private Title"
        document.add_paragraph("Readable document content")
        document.add_picture(str(image))
        document.save(source)
        result = self._clean_and_check(source)
        self.assertEqual(Document(result).paragraphs[0].text, "Readable document content")
        with zipfile.ZipFile(result) as archive:
            for name in archive.namelist():
                self.assertNotIn(b"Private Author", archive.read(name))
                self.assertNotIn(b"Private Editor", archive.read(name))
                self.assertNotIn(b"Private Embedded Author", archive.read(name))

    def test_corrupt_supported_extension_fails_without_publishing(self):
        source = self.home / "private.png"
        source.write_bytes(b"\x89PNG\r\n\x1a\ntruncated")
        with self.assertRaises(cleaner.CleaningError):
            cleaner.clean_file(source, self.home)
        self.assertEqual(list((self.home / "OnionDesk-Clean").iterdir()), [])


if __name__ == "__main__":
    unittest.main()
